"""Exports the knowledge base to a review spreadsheet and applies the completed review.

    python knowledge_base/review_sheet.py export             # writes review_sheet.xlsx
    python knowledge_base/review_sheet.py apply --dry-run    # shows what would change
    python knowledge_base/review_sheet.py apply              # updates sources.json

Decisions: Approve (summary is accurate), Edit (use the corrected summary),
Reject (excerpt is never used). Rows without a decision stay as they are.
Nothing is written if any row has a problem.
"""
import argparse
import json
import re
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

KB_DIR = Path(__file__).resolve().parent
SOURCES_PATH = KB_DIR / "sources.json"
SHEET_PATH = KB_DIR / "review_sheet.xlsx"

COLUMNS = [("Excerpt ID", 22), ("Source", 36), ("URL", 40), ("Section", 30),
           ("Verbatim quote", 60), ("Draft summary", 60), ("Tags", 26), ("Current status", 16),
           ("Reviewer", 18), ("Decision", 12), ("Corrected summary", 60), ("Notes", 40)]
REVIEW_COLUMNS = {"Reviewer", "Decision", "Corrected summary", "Notes"}
DECISIONS = ("Approve", "Edit", "Reject")
RESERVED_NAMES = {"pending", "rejected"}

INSTRUCTIONS = [
    "How to review the knowledge base",
    "",
    "1. Open the URL and find the verbatim quote (python knowledge_base/check_quotes.py confirms every quote is on its page).",
    "2. Read the draft summary. It must say only what the quote and its context say, in plain language.",
    "3. Fill in the yellow columns on the Review sheet:",
    "   Reviewer: your name.",
    "   Decision: Approve if the summary is accurate; Edit if it needs changes (write the new text in Corrected summary); Reject if the excerpt should never be used.",
    "   Notes: optional.",
    "4. Leave Decision blank for rows you have not reviewed; they stay pending.",
    "5. Do not change the quote, ID or other grey columns.",
    "6. Save the file and run: python knowledge_base/review_sheet.py apply --dry-run, then without --dry-run.",
]


def load_sources(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def save_sources(data, path):
    Path(path).write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def squash(text):
    return re.sub(r"\s+", " ", str(text or "")).strip()


def export(sources_path=SOURCES_PATH, sheet_path=SHEET_PATH):
    data = load_sources(sources_path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Review"
    ws.append([name for name, _ in COLUMNS])

    for source in data["sources"]:
        for e in source["excerpts"]:
            review = e.get("review", {})
            edited = review.get("decision") == "Edit"
            ws.append([
                e["excerpt_id"], f"{source['title']} ({source['publisher']}, {source['year']})",
                source["url"], e["section"], e["quote"],
                review.get("original_summary", e["summary"]), ", ".join(e["tags"]), e["verified_by"],
                review.get("reviewer", ""), review.get("decision", ""),
                e["summary"] if edited else "", review.get("notes", ""),
            ])

    editable = PatternFill("solid", fgColor="FFF2CC")
    locked = PatternFill("solid", fgColor="EDEDED")
    for col, (name, width) in enumerate(COLUMNS, start=1):
        letter = ws.cell(row=1, column=col).column_letter
        ws.column_dimensions[letter].width = width
        header = ws.cell(row=1, column=col)
        header.font = Font(bold=True)
        header.fill = editable if name in REVIEW_COLUMNS else locked
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    decision_col = [name for name, _ in COLUMNS].index("Decision") + 1
    letter = ws.cell(row=1, column=decision_col).column_letter
    validation = DataValidation(type="list", formula1='"Approve,Edit,Reject"', allow_blank=True)
    validation.add(f"{letter}2:{letter}{ws.max_row}")
    ws.add_data_validation(validation)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    notes = wb.create_sheet("Instructions")
    for line in INSTRUCTIONS:
        notes.append([line])
    notes["A1"].font = Font(bold=True)
    notes.column_dimensions["A"].width = 140

    wb.save(sheet_path)
    return ws.max_row - 1


def read_rows(sheet_path):
    ws = load_workbook(sheet_path, data_only=True)["Review"]
    header = [c.value for c in ws[1]]
    expected = [name for name, _ in COLUMNS]
    if header[:len(expected)] != expected:
        raise ValueError(f"Unexpected columns in the Review sheet; expected {expected}")
    for number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        row = {name: squash(v) for name, v in zip(expected, values)}
        if row["Excerpt ID"]:
            yield number, row


def apply(sheet_path=SHEET_PATH, sources_path=SOURCES_PATH, dry_run=False, today=None):
    """Returns (changes, errors). sources.json is only written when there are no errors."""
    today = today or date.today().isoformat()
    data = load_sources(sources_path)
    excerpts = {e["excerpt_id"]: e for s in data["sources"] for e in s["excerpts"]}
    changes, errors, seen = [], [], set()

    for number, row in read_rows(sheet_path):
        excerpt_id = row["Excerpt ID"]
        where = f"row {number} ({excerpt_id})"
        excerpt = excerpts.get(excerpt_id)
        if excerpt is None:
            errors.append(f"{where}: unknown excerpt ID")
            continue
        if excerpt_id in seen:
            errors.append(f"{where}: excerpt appears more than once")
            continue
        seen.add(excerpt_id)
        if row["Verbatim quote"] != squash(excerpt["quote"]):
            errors.append(f"{where}: the quote was changed; quotes must match the source word for word")
            continue
        if not row["Decision"]:
            continue

        decision = row["Decision"].capitalize()
        reviewer, corrected = row["Reviewer"], row["Corrected summary"]
        if decision not in DECISIONS:
            errors.append(f"{where}: Decision must be Approve, Edit or Reject, not '{row['Decision']}'")
            continue
        if not reviewer or reviewer.lower() in RESERVED_NAMES:
            errors.append(f"{where}: a reviewer name is required")
            continue
        if decision == "Edit" and not corrected:
            errors.append(f"{where}: Edit needs a corrected summary")
            continue
        if decision != "Edit" and corrected:
            errors.append(f"{where}: Corrected summary is only used with Edit")
            continue

        previous = excerpt.get("review", {})
        original = previous.get("original_summary", excerpt["summary"])
        review = {"reviewer": reviewer, "decision": decision, "date": previous.get("date", today)}
        summary = original
        if decision == "Edit":
            summary = corrected
            review["original_summary"] = original
        if row["Notes"]:
            review["notes"] = row["Notes"]
        verified_by = "rejected" if decision == "Reject" else reviewer

        unchanged = (excerpt["verified_by"] == verified_by and excerpt["summary"] == summary
                     and {k: v for k, v in previous.items() if k != "date"}
                     == {k: v for k, v in review.items() if k != "date"})
        if unchanged:
            continue
        review["date"] = today
        changes.append(f"{excerpt_id}: {excerpt['verified_by']} -> {verified_by}"
                       + (" (summary edited)" if summary != excerpt["summary"] else ""))
        excerpt.update(verified_by=verified_by, summary=summary, review=review)

    if not errors and changes and not dry_run:
        save_sources(data, sources_path)
    return changes, errors


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("command", choices=["export", "apply"])
    parser.add_argument("--dry-run", action="store_true", help="show changes without writing sources.json")
    args = parser.parse_args()

    if args.command == "export":
        print(f"Wrote {export()} excerpts to {SHEET_PATH.name}")
        return 0

    changes, errors = apply(dry_run=args.dry_run)
    for line in errors:
        print(f"ERROR {line}")
    for line in changes:
        print(line)
    if errors:
        print(f"\n{len(errors)} problem(s); sources.json was not changed")
        return 1
    if not changes:
        print("No changes")
    elif args.dry_run:
        print(f"\n{len(changes)} change(s) would be made (dry run, nothing written)")
    else:
        print(f"\nUpdated {len(changes)} excerpt(s) in {SOURCES_PATH.name}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
