import json
import shutil

import pytest
from openpyxl import load_workbook

from knowledge_base import review_sheet
from knowledge_base.review_sheet import COLUMNS, apply, export

COL = {name: i for i, (name, _) in enumerate(COLUMNS, start=1)}


@pytest.fixture
def workspace(tmp_path):
    sources = tmp_path / "sources.json"
    shutil.copy(review_sheet.SOURCES_PATH, sources)
    sheet = tmp_path / "review_sheet.xlsx"
    export(sources, sheet)
    return sources, sheet


def fill(sheet, edits):
    """edits: {excerpt_id: {column name: value}}"""
    wb = load_workbook(sheet)
    ws = wb["Review"]
    for row in range(2, ws.max_row + 1):
        excerpt_id = ws.cell(row=row, column=COL["Excerpt ID"]).value
        for column, value in edits.get(excerpt_id, {}).items():
            ws.cell(row=row, column=COL[column]).value = value
    wb.save(sheet)


def excerpt(sources, excerpt_id):
    data = json.loads(sources.read_text(encoding="utf-8"))
    return next(e for s in data["sources"] for e in s["excerpts"] if e["excerpt_id"] == excerpt_id)


def test_export_has_one_row_per_excerpt_and_blank_review_columns(workspace):
    sources, sheet = workspace
    data = json.loads(sources.read_text(encoding="utf-8"))
    total = sum(len(s["excerpts"]) for s in data["sources"])
    ws = load_workbook(sheet)["Review"]
    assert [c.value for c in ws[1]] == [name for name, _ in COLUMNS]
    assert ws.max_row - 1 == total
    for row in ws.iter_rows(min_row=2, values_only=True):
        assert row[COL["Current status"] - 1] == "pending"
        assert all(row[COL[c] - 1] in (None, "") for c in review_sheet.REVIEW_COLUMNS)
    validations = ws.data_validations.dataValidation
    assert any('"Approve,Edit,Reject"' == v.formula1 for v in validations)


def test_unreviewed_sheet_changes_nothing(workspace):
    sources, sheet = workspace
    before = sources.read_text(encoding="utf-8")
    assert apply(sheet, sources) == ([], [])
    assert sources.read_text(encoding="utf-8") == before


def test_approve_edit_and_reject(workspace):
    sources, sheet = workspace
    original = excerpt(sources, "CMS-HRRP-02")["summary"]
    fill(sheet, {
        "CMS-HRRP-01": {"Reviewer": "Dr. Rivera", "Decision": "Approve"},
        "CMS-HRRP-02": {"Reviewer": "Dr. Rivera", "Decision": "edit",
                        "Corrected summary": "Medicare payment cuts for excess readmissions are capped at 3%.",
                        "Notes": "Clarified wording"},
        "STRACK-2014-03": {"Reviewer": "A. Student", "Decision": "Reject"},
    })
    changes, errors = apply(sheet, sources, today="2026-10-20")
    assert errors == [] and len(changes) == 3

    approved = excerpt(sources, "CMS-HRRP-01")
    assert approved["verified_by"] == "Dr. Rivera"
    assert approved["review"] == {"reviewer": "Dr. Rivera", "decision": "Approve", "date": "2026-10-20"}

    edited = excerpt(sources, "CMS-HRRP-02")
    assert edited["verified_by"] == "Dr. Rivera"
    assert edited["summary"] == "Medicare payment cuts for excess readmissions are capped at 3%."
    assert edited["review"]["original_summary"] == original
    assert edited["review"]["notes"] == "Clarified wording"

    assert excerpt(sources, "STRACK-2014-03")["verified_by"] == "rejected"
    assert excerpt(sources, "STRACK-2014-01")["verified_by"] == "pending"


def test_reapplying_and_reexporting_is_stable(workspace):
    sources, sheet = workspace
    fill(sheet, {"CMS-HRRP-02": {"Reviewer": "R", "Decision": "Edit", "Corrected summary": "New text."}})
    apply(sheet, sources, today="2026-10-20")
    assert apply(sheet, sources, today="2026-10-21") == ([], [])

    # A fresh export shows the original draft and the correction side by side
    export(sources, sheet)
    ws = load_workbook(sheet)["Review"]
    row = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] == "CMS-HRRP-02")
    assert row[COL["Corrected summary"] - 1] == "New text."
    assert row[COL["Draft summary"] - 1] != "New text."
    assert apply(sheet, sources, today="2026-10-22") == ([], [])
    assert excerpt(sources, "CMS-HRRP-02")["review"]["date"] == "2026-10-20"


@pytest.mark.parametrize("edits,message", [
    ({"Decision": "Approve"}, "a reviewer name is required"),
    ({"Reviewer": "pending", "Decision": "Approve"}, "a reviewer name is required"),
    ({"Reviewer": "R", "Decision": "Maybe"}, "Decision must be Approve, Edit or Reject"),
    ({"Reviewer": "R", "Decision": "Edit"}, "Edit needs a corrected summary"),
    ({"Reviewer": "R", "Decision": "Approve", "Corrected summary": "x"}, "only used with Edit"),
    ({"Verbatim quote": "Something else", "Reviewer": "R", "Decision": "Approve"}, "quote was changed"),
])
def test_problems_block_the_whole_update(workspace, edits, message):
    sources, sheet = workspace
    before = sources.read_text(encoding="utf-8")
    fill(sheet, {"CMS-HRRP-01": edits, "CMS-HRRP-02": {"Reviewer": "R", "Decision": "Approve"}})
    changes, errors = apply(sheet, sources)
    assert len(errors) == 1 and message in errors[0] and "CMS-HRRP-01" in errors[0]
    assert sources.read_text(encoding="utf-8") == before


def test_dry_run_writes_nothing(workspace):
    sources, sheet = workspace
    before = sources.read_text(encoding="utf-8")
    fill(sheet, {"CMS-HRRP-01": {"Reviewer": "R", "Decision": "Approve"}})
    changes, errors = apply(sheet, sources, dry_run=True)
    assert len(changes) == 1 and not errors
    assert sources.read_text(encoding="utf-8") == before


def test_committed_sheet_matches_the_knowledge_base():
    rows = {row["Excerpt ID"]: row for _, row in review_sheet.read_rows(review_sheet.SHEET_PATH)}
    data = json.loads(review_sheet.SOURCES_PATH.read_text(encoding="utf-8"))
    ids = {e["excerpt_id"] for s in data["sources"] for e in s["excerpts"]}
    assert set(rows) == ids
